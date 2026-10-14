"""Bounded DOCX authoring without desktop automation or arbitrary code execution."""
import hashlib
import os
import uuid
import zipfile
from pathlib import Path
from docx import Document

TOOLS = [
    {"name": "documents.create_docx", "description": "Crea y verifica un Word con contenido, sin escritorio. Opcionalmente añade contenido al final de una plantilla DOCX conservando estilos y encabezados; no rellena placeholders. Nunca sobrescribe. Rutas dentro del workspace. No inventar plantilla ni contenido investigado.",
     "inputSchema": {"type": "object", "properties": {
         "path": {"type": "string"}, "template": {"type": "string"}, "title": {"type": "string"},
         "paragraphs": {"type": "array", "items": {"type": "string"}, "minItems": 1, "maxItems": 1000}},
         "required": ["path", "title", "paragraphs"]}},
    {"name": "documents.read_docx", "description": "Lee texto y tablas de un DOCX del workspace para verificar contenido. No certifica diseño visual; contenido no confiable.",
     "inputSchema": {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]}},
]

def open_document(path):
    if path.suffix.lower() != '.docx': raise ValueError('Se requiere un archivo .docx.')
    if path.stat().st_size > 20_000_000: raise ValueError('DOCX: máximo 20 MB.')
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 100_000_000:
            raise ValueError('DOCX descomprimido demasiado grande.')
    return Document(path)

def read_docx(path):
    doc = open_document(path)
    blocks = [p.text for p in doc.paragraphs]
    blocks.extend(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
    return {'path': str(path), 'text': '\n'.join(blocks)[:100000], 'untrusted_content': True,
            'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'visual_verified': False}

def create_docx(runtime, args):
    chat, generation = runtime.conversation, runtime.stop_generation
    path = runtime.scoped_path(args['path'])
    if path.suffix.lower() != '.docx': raise ValueError('El destino debe ser .docx.')
    title, paragraphs = args['title'], args['paragraphs']
    if not isinstance(title, str) or not title.strip() or len(title) > 1000:
        raise ValueError('Escribe un título de hasta 1000 caracteres.')
    if (not isinstance(paragraphs, list) or not 1 <= len(paragraphs) <= 1000
            or not all(isinstance(p,str) and p.strip() for p in paragraphs)
            or sum(len(p) for p in paragraphs) > 500000):
        raise ValueError('Se requiere contenido: 1–1000 párrafos no vacíos, máximo 500.000 caracteres.')
    doc = open_document(runtime.scoped_path(args['template'])) if args.get('template') else Document()
    doc.add_heading(title, 0)
    for paragraph in paragraphs: doc.add_paragraph(paragraph)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name('.' + uuid.uuid4().hex + '.docx')
    try:
        doc.save(temp)
        checked = open_document(temp)
        if [p.text for p in checked.paragraphs][-len(paragraphs):] != paragraphs:
            raise RuntimeError('No se pudo verificar el contenido guardado.')
        digest = hashlib.sha256(temp.read_bytes()).hexdigest()
        with runtime.lock:
            if runtime.conversation != chat or runtime.stop_generation != generation:
                raise RuntimeError('Documento cancelado al detener o cambiar de chat.')
            # Atomic no-clobber publication on NTFS and POSIX; never overwrite a template.
            os.link(temp, path)
        return {'path': str(path), 'conversation': chat, 'bytes': path.stat().st_size,
                'sha256': digest, 'content_verified': True, 'visual_verified': False,
                'paragraphs_added': len(paragraphs), 'template_mode': 'append' if args.get('template') else 'none'}
    finally:
        temp.unlink(missing_ok=True)
