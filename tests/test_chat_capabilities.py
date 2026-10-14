import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from arise_app.assistant import Assistant

class ChatCapabilitiesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.r = Assistant(self.root/'data')
        self.work = self.root/'work'; self.work.mkdir()
        self.r.select_project(str(self.work))
        self.r.storage.config.update(pi_command=[sys.executable,str(Path(__file__).parent/'fake_agent.py'),'--pi'],mcp={})
    def tearDown(self):
        self.r.close(); self.r.storage.db.close(); self.tmp.cleanup()
    def test_two_messages_voice_settings_and_stop_retain_process(self):
        r=self.r
        r.prompt('primero'); self.assertTrue(r.settled.wait(5)); process=r.pi.process
        r.prompt('segundo'); self.assertTrue(r.settled.wait(5))
        r.settings({'forge_rate':'+6%'})
        self.assertIs(r.pi.process,process)
        r.stop(); self.assertIs(r.pi.process,process)
        self.assertFalse(r.desktop); self.assertTrue(r.desktop_halted)
        self.assertEqual(len([m for m in r.storage.messages(r.conversation) if m['role']=='user']),2)
    def test_permissions_do_not_copy_and_late_events_do_not_cross_chat(self):
        r=self.r; first=r.conversation
        r.storage.config['code_enabled']=True
        r.connect_pi(); epoch=r.pi_epoch
        second=r.switch_conversation()['id']
        self.assertFalse(r.storage.config['code_enabled'])
        r.session_event(epoch,{'type':'message_end','message':{'role':'assistant','content':[{'type':'text','text':'LATE'}]}})
        self.assertEqual(r.storage.messages(second),[])
        r.switch_conversation(first)
        self.assertTrue(r.storage.config['code_enabled'])
        self.assertFalse(r.desktop)
    def test_real_docx_template_content_readback_and_no_overwrite(self):
        doc=Document(); doc.sections[0].header.paragraphs[0].text='WEG — plantilla de prueba'
        doc.save(self.work/'template.docx')
        args={'path':'result.docx','template':'template.docx','title':'Beneficios E3',
              'paragraphs':['Contenido de prueba con acentos: Colombia.','Segundo párrafo verificable.']}
        result=json.loads(self.r.call_tool('documents.create_docx',args)['content'][0]['text'])
        self.assertTrue(result['content_verified']); self.assertFalse(result['visual_verified'])
        actual=Document(self.work/'result.docx')
        self.assertEqual(actual.sections[0].header.paragraphs[0].text,'WEG — plantilla de prueba')
        self.assertEqual([p.text for p in actual.paragraphs][-2:],args['paragraphs'])
        self.assertIn('Colombia',self.r.call_tool('documents.read_docx',{'path':'result.docx'})['content'][0]['text'])
        with self.assertRaises(FileExistsError): self.r.call_tool('documents.create_docx',args)
        with self.assertRaises(ValueError): self.r.call_tool('documents.create_docx',{**args,'path':'../outside.docx'})
        with self.assertRaises(ValueError): self.r.call_tool('documents.create_docx',{**args,'paragraphs':[]})
    def test_document_cancel_during_generation_never_publishes(self):
        import arise_app.documents as documents
        original=documents.open_document
        def stopped(path):
            result=original(path); self.r.stop(); return result
        with patch.object(documents,'open_document',side_effect=stopped):
            with self.assertRaisesRegex(RuntimeError,'cancelado'):
                self.r.call_tool('documents.create_docx',{'path':'cancelled.docx','title':'T','paragraphs':['Contenido']})
        self.assertFalse((self.work/'cancelled.docx').exists())
    def test_stop_after_delegation_closes_owned_tree_without_auto_restart(self):
        r=self.r; r.connect_pi(); process=r.pi.process
        r.pi_event({'type':'tool_execution_start','toolName':'subagent_run'})
        r.stop()
        self.assertIsNotNone(process.poll()); self.assertIsNone(r.pi)
        self.assertTrue(r.desktop_halted)
